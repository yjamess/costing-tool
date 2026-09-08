"""Load Lustre MASTER rates from the RatesMaster workbook.

The LLM must never invent prices. Python looks up Item IDs (Ref) here
and applies margin rules from the workbook READ ME:

- LUSTRE FEE / LUSTRE EXP → sell = net (no uplift)
- THIRD PARTY → sell = net * (1 + margin)

openpyxl docs: https://openpyxl.readthedocs.io/
"""

from pathlib import Path

from openpyxl import load_workbook
from pydantic import BaseModel, Field


DEFAULT_RATES_PATH = Path("docs/Lustre_RatesMaster_v1.0.xlsx")

# Header row on the MASTER tab (1-based, as in Excel).
_HEADER_ROW = 4
_DATA_START_ROW = 5


class RateLine(BaseModel):
    """One row from the MASTER rate library."""

    ref: str = Field(description="Item ID, e.g. C.cp200_0")
    section_code: str
    section: str
    group: str
    item_id: str
    description: str
    unit: str
    cost_class: str
    supplier: str | None = None
    net_rate: float
    basis: str | None = None
    sell_10: float | None = None
    sell_15: float | None = None
    sell_20: float | None = None


def load_master_rates(
    workbook_path: str | Path = DEFAULT_RATES_PATH,
) -> dict[str, RateLine]:
    """Read the MASTER sheet into a dict keyed by Ref.

    Args:
        workbook_path: Path to Lustre_RatesMaster_v1.0.xlsx

    Returns:
        Mapping of ref → RateLine

    Raises:
        FileNotFoundError: If the workbook is missing.
        ValueError: If MASTER is missing or no rate rows are found.
    """
    path = Path(workbook_path)
    if not path.is_file():
        raise FileNotFoundError(f"Rates workbook not found: {path}")

    # data_only=True uses cached calculated values for formula cells.
    wb = load_workbook(path, data_only=True, read_only=True)
    if "MASTER" not in wb.sheetnames:
        raise ValueError(f"No MASTER sheet in {path.name}. Found: {wb.sheetnames}")

    ws = wb["MASTER"]
    rates: dict[str, RateLine] = {}

    for row in ws.iter_rows(min_row=_DATA_START_ROW, values_only=True):
        ref = row[0]
        if not ref or not isinstance(ref, str):
            continue

        net = row[9]  # column J — Net rate £
        if net is None:
            continue

        line = RateLine(
            ref=ref.strip(),
            section_code=str(row[1] or ""),
            section=str(row[2] or ""),
            group=str(row[3] or ""),
            item_id=str(row[4] or ""),
            description=str(row[5] or ""),
            unit=str(row[6] or ""),
            cost_class=str(row[7] or "").strip().upper(),
            supplier=str(row[8]) if row[8] is not None else None,
            net_rate=float(net),
            basis=str(row[10]) if row[10] is not None else None,
            sell_10=float(row[12]) if row[12] is not None else None,
            sell_15=float(row[13]) if row[13] is not None else None,
            sell_20=float(row[14]) if row[14] is not None else None,
        )
        rates[line.ref] = line

    wb.close()

    if not rates:
        raise ValueError(f"No rate lines loaded from MASTER in {path.name}")

    return rates


def sell_rate(line: RateLine, margin_pct: int = 15) -> float:
    """Return the sell rate for a line at 10, 15 or 20% job margin.

    Prefer the pre-computed Sell columns from MASTER when present.
    Otherwise apply the READ ME rule in Python.
    """
    if margin_pct not in (10, 15, 20):
        raise ValueError("margin_pct must be 10, 15 or 20")

    precomputed = {10: line.sell_10, 15: line.sell_15, 20: line.sell_20}[margin_pct]
    if precomputed is not None:
        return precomputed

    if line.cost_class in {"LUSTRE FEE", "LUSTRE EXP"}:
        return line.net_rate

    if line.cost_class == "THIRD PARTY":
        return round(line.net_rate * (1 + margin_pct / 100), 2)

    raise ValueError(f"Unknown cost_class for {line.ref}: {line.cost_class!r}")


if __name__ == "__main__":
    # Quick smoke test: uv run python rates.py
    catalogue = load_master_rates()
    print(f"Loaded {len(catalogue)} MASTER lines")

    sample_refs = ["A.sup_full", "D.ws_day", "C.cp200_0", "I.si_lustre"]
    for ref in sample_refs:
        line = catalogue[ref]
        print(
            f"{line.ref}: {line.description[:50]}... | "
            f"{line.unit} | {line.cost_class} | "
            f"net={line.net_rate} sell@15={sell_rate(line, 15)}"
        )
